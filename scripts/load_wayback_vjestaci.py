"""Load MOJ Stalni sudski vještaci 2017 XLSX (recovered from Wayback Machine)
into Toronto's expert_witnesses table.

Source: pravosudje.gov.hr 'Stalni sudski vjestaci_17 1 2017.xlsx', snapshot
2018-01-20.

Strategy: dedupe against existing experts by normalized name. Existing rows
get expert_type='vjestak'. New rows inserted with that type + structured
fields.
"""
import openpyxl, re, subprocess, sys, unicodedata

XLSX = '/tmp/vjestaci_2017.xlsx'

def ssh_psql_select(sql):
    cmd = ['ssh', 'toronto-sudacka',
           f"docker exec sudacka-mreza-db-1 psql -U postgres -d sudacka_mreza -At -F'|' -c \"{sql}\""]
    return subprocess.check_output(cmd, text=True)

def ssh_psql_exec(sql):
    cmd = ['ssh', 'toronto-sudacka', 'docker exec -i sudacka-mreza-db-1 psql -U postgres -d sudacka_mreza -At']
    r = subprocess.run(cmd, input=sql, capture_output=True, text=True, check=False)
    if r.returncode != 0:
        print('STDERR:', r.stderr[:500], file=sys.stderr)
    return r.stdout, r.returncode

def norm_name(s):
    s = unicodedata.normalize('NFKD', s or '').encode('ascii','ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()

def slugify(s):
    s = unicodedata.normalize('NFKD', s or '').encode('ascii','ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-')[:64] or 'vjestak'

def split_areas(s):
    if not s: return []
    s = s.replace('_x000D_', ' ').replace('\r', ' ').replace('\n', ' ')
    parts = re.split(r'\s*(?:,|;|/|-\s+)\s*', s)
    return [p.strip(' -•·.') for p in parts if p.strip(' -•·.') and len(p.strip()) > 2]

def sql_quote(s):
    if s is None: return 'NULL'
    return "'" + str(s).replace("'", "''") + "'"

# Existing experts for dedupe
print('[1/3] load existing experts')
out = ssh_psql_select("SELECT id, name FROM expert_witnesses ORDER BY id")
existing = {}
for line in out.splitlines():
    if '|' not in line: continue
    eid, name = line.split('|', 1)
    existing[norm_name(name)] = int(eid)
print(f'   {len(existing)} existing experts')

# Parse XLSX
print('[2/3] parse XLSX')
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(min_row=2, values_only=True))
print(f'   {len(rows)} rows in XLSX')

planned = []
for r in rows:
    first = (r[0] or '').strip() if r[0] else ''
    last = (r[1] or '').strip() if r[1] else ''
    court_code = (r[2] or '').strip() if r[2] else ''
    city = (r[3] or '').strip() if r[3] else ''
    address = (r[4] or '').strip() if r[4] else ''
    areas_str = (r[5] or '').strip() if r[5] else ''
    if not first and not last: continue
    name = f'{first} {last}'.strip()
    nk = norm_name(name)
    if not nk: continue
    planned.append({
        'name': name, 'nk': nk,
        'court_code': court_code, 'city': city, 'address': address,
        'areas': split_areas(areas_str),
    })

new_inserts = [p for p in planned if p['nk'] not in existing]
existing_updates = [p for p in planned if p['nk'] in existing]
print(f'   {len(planned)} parsed, {len(existing_updates)} match existing, {len(new_inserts)} new')

# Build SQL
print('[3/3] apply')
batches = []
# 1. tag matched existing as expert_type='vjestak' if not already set
for p in existing_updates:
    eid = existing[p['nk']]
    batches.append(f"UPDATE expert_witnesses SET expert_type = 'vjestak' WHERE id = {eid} AND (expert_type IS NULL OR expert_type = '');")
# 2. insert new
for p in new_inserts:
    name_q = sql_quote(p['name'])
    address_q = sql_quote(p['address'] or None)
    city_q = sql_quote(p['city'] or None)
    slug = f"{slugify(p['name'])}-wb2017"
    slug_q = sql_quote(slug)
    batches.append(
        f"INSERT INTO expert_witnesses (name, expert_type, address, city, lang, slug, verified, created_at, updated_at) "
        f"VALUES ({name_q}, 'vjestak', {address_q}, {city_q}, 'hr', {slug_q}, false, NOW(), NOW()) "
        f"ON CONFLICT (slug) DO NOTHING;"
    )

print(f'   {len(batches)} statements; applying…')
# Apply in chunks of 500 to avoid huge stdin
CHUNK = 500
errs = 0
for i in range(0, len(batches), CHUNK):
    out, code = ssh_psql_exec('BEGIN;\n' + '\n'.join(batches[i:i+CHUNK]) + '\nCOMMIT;')
    if code != 0: errs += 1
print(f'   chunks done; errors={errs}')

# Status
out2 = ssh_psql_select("SELECT COALESCE(expert_type, '(null)') AS t, COUNT(*) FROM expert_witnesses GROUP BY 1 ORDER BY 2 DESC;")
print('---post---')
print(out2)
