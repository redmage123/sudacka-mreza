"""Materialize specialty areas for the wayback-loaded vještaci.

Source: column 6 ('Područje vještačenja/procjene') of the 2017 MOJ XLSX.

For each row, splits the area cell into individual entries and inserts them
into expert_witnesses_speciality_areas. Resolves the expert by slug
(<slug>-wb2017) or by normalized name. Skips rows that already have any
speciality rows so we don't double up if rerun.
"""
import openpyxl, re, subprocess, sys, unicodedata

XLSX = '/tmp/vjestaci_2017.xlsx'

def ssh_psql_select(sql):
    cmd = ['ssh', 'toronto-sudacka',
           f"docker exec sudacka-mreza-db-1 psql -U postgres -d sudacka_mreza -At -F'|' -c \"{sql}\""]
    return subprocess.check_output(cmd, text=True)

def ssh_psql_exec(sql):
    cmd = ['ssh', 'toronto-sudacka', 'docker exec -i sudacka-mreza-db-1 psql -U postgres -d sudacka_mreza -At']
    r = subprocess.run(cmd, input=sql, capture_output=True, text=True, check=False)
    if r.returncode != 0:
        print('STDERR:', r.stderr[:500], file=sys.stderr)
    return r.stdout, r.returncode

def deaccent(s):
    return unicodedata.normalize('NFKD', s or '').encode('ascii','ignore').decode().lower()

def norm_name(s):
    return re.sub(r'[^a-z0-9]+', ' ', deaccent(s)).strip()

def slugify(s):
    s = deaccent(s)
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-')[:64] or 'vjestak'

def sql_quote(s):
    if s is None: return 'NULL'
    return "'" + str(s).replace("'", "''") + "'"

def split_areas(s):
    """Split a 'Područje vještačenja' cell. The 2017 XLSX uses comma as
    the primary separator: 'financije, knjigovodstvo' = two entries,
    'proc. polj. zem., procj. nekretnina, procjena zemljišta' = three.
    Some cells use ';' too; handle both."""
    if not s: return []
    s = s.replace('_x000D_', ' ').replace('\r', ' ').replace('\n', ' ')
    # Normalize semicolons to commas so we can split once.
    s = s.replace(';', ',')
    parts = [p.strip(' -•·.') for p in s.split(',')]
    return [p for p in parts if p and len(p) > 1]

# Load XLSX
print('[1/4] parse XLSX')
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    first = (r[0] or '').strip() if r[0] else ''
    last = (r[1] or '').strip() if r[1] else ''
    areas_str = (r[5] or '').strip() if r[5] else ''
    if not (first or last): continue
    name = f'{first} {last}'.strip()
    areas = split_areas(areas_str)
    if areas:
        rows.append((name, areas))
print(f'   {len(rows)} rows with at least one area; total area entries: {sum(len(a) for _, a in rows)}')

# Map expert ids
print('[2/4] resolve expert ids')
out = ssh_psql_select("SELECT id, slug, name FROM expert_witnesses WHERE expert_type = 'vjestak'")
by_slug, by_name = {}, {}
for line in out.splitlines():
    if '|' not in line: continue
    parts = line.split('|', 2)
    if len(parts) < 3: continue
    eid, slug, name = int(parts[0]), parts[1], parts[2]
    by_slug[slug] = eid
    by_name[norm_name(name)] = eid
print(f'   {len(by_slug)} vještak experts by slug')

# Existing speciality counts per parent — skip rows that already have any
# entries (we don't want to double-up on a rerun).
print('[3/4] load existing speciality counts')
out = ssh_psql_select("SELECT _parent_id, COUNT(*) FROM expert_witnesses_speciality_areas GROUP BY _parent_id")
have_existing = set()
for line in out.splitlines():
    if '|' not in line: continue
    pid, _ = line.split('|', 1)
    have_existing.add(int(pid))
print(f'   {len(have_existing)} experts already have at least one area row')

# Build inserts
print('[4/4] build inserts')
batches = []
linked_experts = 0
skipped_no_expert = 0
skipped_existing = 0
total_area_rels = 0

for name, areas in rows:
    slug = f'{slugify(name)}-wb2017'
    eid = by_slug.get(slug) or by_name.get(norm_name(name))
    if eid is None:
        skipped_no_expert += 1
        continue
    if eid in have_existing:
        skipped_existing += 1
        continue
    linked_experts += 1
    for ord_i, area in enumerate(areas, start=1):
        # id must be unique varchar; combine eid + order.
        row_id = f'wb2017-{eid}-{ord_i}'
        batches.append(
            f"INSERT INTO expert_witnesses_speciality_areas (_order, _parent_id, id, area, sub_area) "
            f"VALUES ({ord_i}, {eid}, {sql_quote(row_id)}, {sql_quote(area)}, NULL) "
            f"ON CONFLICT (id) DO NOTHING;"
        )
        total_area_rels += 1

print(f'   plan: {linked_experts} experts, {total_area_rels} area rows; '
      f'skipped {skipped_no_expert} unmatched, {skipped_existing} already populated')

if '--dry-run' in sys.argv:
    print('dry-run; not writing')
    sys.exit(0)

CHUNK = 500
errs = 0
for i in range(0, len(batches), CHUNK):
    _, code = ssh_psql_exec('BEGIN;\n' + '\n'.join(batches[i:i+CHUNK]) + '\nCOMMIT;')
    if code != 0: errs += 1
print(f'   applied {len(batches)} inserts; errors={errs}')

out = ssh_psql_select("SELECT COUNT(*) FROM expert_witnesses_speciality_areas;")
print(f"   total speciality area rows in DB: {out.strip()}")
out = ssh_psql_select("SELECT COUNT(DISTINCT _parent_id) FROM expert_witnesses_speciality_areas;")
print(f"   distinct experts with >=1 area: {out.strip()}")
